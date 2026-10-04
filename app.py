"""Sistem Inventory Stok Grade - Flask + SQLite + Excel (portable).

Model data ternormalisasi: transaksi hanya menyimpan ID master. Motif, Jenis, Rumus,
Jumlah produksi dihitung saat ditampilkan, jadi tidak ada data ganda yang bisa berbeda.
"""
import os, sys, io, re, json, shutil, socket, sqlite3, threading, webbrowser
import datetime as dt
from flask import Flask, g, request, jsonify, render_template, send_file, redirect, abort
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FROZEN = getattr(sys, "frozen", False)
BASE = os.path.dirname(sys.executable) if FROZEN else os.path.dirname(os.path.abspath(__file__))
RES = getattr(sys, "_MEIPASS", BASE)
DATA_DIR = os.path.join(BASE, "data")
BACKUP_DIR = os.path.join(BASE, "backup")
DB_PATH = os.path.join(DATA_DIR, "inventory.db")

app = Flask(__name__, template_folder=os.path.join(RES, "templates"),
            static_folder=os.path.join(RES, "static"))

# --------------------------------------------------------------------------
# Definisi tabel master + produksi
# --------------------------------------------------------------------------
CATS = ["wadimor", "junior", "celup", "biasa", "perbaikan", "neci",
        "rusak", "jahit", "melipat", "cabut", "washing"]
JUMLAH_EXPR = "(" + "+".join(f"t.{c}" for c in CATS) + ")"

def C(n, l, t="text", **k):
    return dict(n=n, l=l, t=t, **k)

REFS = {  # sumber dropdown: (tabel, kolom nilai)
    "jenis": ("ref_jenis", "nama"), "ket": ("ref_ket", "nama"), "dept": ("ref_dept", "nama"),
    "ketprod": ("ref_ketprod", "nama"), "pengrajin": ("ref_pengrajin", "nama"),
    "karyawan": ("ref_karyawan", "nama"), "motif": ("ref_motif", "kode_motif"),
}
# tabel/kolom yang memakai master (untuk mencegah hapus master yang masih dipakai)
USAGE = {
    "jenis": [("ref_motif", "jenis_id")],
    "ket": [("dokumen_item", "ket_id")], "dept": [("dokumen", "dept_id")],
    "ketprod": [("produksi", "ket_id")], "pengrajin": [("dokumen", "pengrajin_id")],
    "karyawan": [("produksi", "nama_id"), ("produksi", "pasangan_id")],
    "motif": [("produksi", "kode_motif_id"), ("dokumen_item", "motif_id"), ("opening", "kode_motif_id")],
}

SPECS = {
    "ket": dict(table="ref_ket", title="Ket & Rumus", color="70AD47",
                cols=[C("nama", "Nama Ket", req=1, uniq=1), C("kode", "Kode Rumus", req=1)]),
    "dept": dict(table="ref_dept", title="Dept & Rumus", color="ED7D31",
                 cols=[C("nama", "Dept", req=1, uniq=1), C("kode", "Rumus", req=1)]),
    "ketprod": dict(table="ref_ketprod", title="Ket Produksi", color="7030A0",
                    cols=[C("nama", "Ket Produksi", req=1, uniq=1)]),
    "pengrajin": dict(table="ref_pengrajin", title="Pengrajin", color="FFC000",
                      cols=[C("nama", "Pengrajin", req=1, uniq=1)]),
    "karyawan": dict(table="ref_karyawan", title="Karyawan", color="8EA9DB",
                     cols=[C("nik", "NIK", uniq=1), C("nama", "Nama", req=1, uniq=1)]),
    "jenis": dict(table="ref_jenis", title="Jenis", color="808080",
                  cols=[C("nama", "Jenis", req=1, uniq=1, upper=1)]),
    "motif": dict(table="ref_motif", title="Motif", color="4472C4",
                  cols=[C("kode_motif", "Kode Motif", req=1, uniq=1),
                        C("nama_motif", "Nama Motif", req=1),
                        C("jenis", "Jenis", "ref", src="jenis", req=1)]),
    "opening": dict(table="opening", title="Stok Awal Sistem", color="808080",
                    joins=["LEFT JOIN ref_jenis rj ON rj.id=r_kode_motif.jenis_id"],
                    cols=[C("kode_motif", "Kode Motif", "ref", src="motif", req=1, uniq=1),
                          C("motif", "Motif", "auto", sql="r_kode_motif.nama_motif", hc="FF0000"),
                          C("jenis", "Jenis", "auto", sql="rj.nama", hc="FF0000"),
                          C("jumlah", "Jumlah", "int", req=1),
                          C("tanggal_mulai", "Berlaku Mulai Tanggal", "date")]),
    "produksi": dict(table="produksi", title="Produksi", color="4472C4", datecol="tgl_produksi",
                     joins=["LEFT JOIN ref_jenis rj ON rj.id=r_kode_motif.jenis_id"],
                     cols=[C("tgl_produksi", "Tanggal Produksi", "date", req=1),
                           C("tgl_masuk_sanggan", "Tanggal Masuk Sanggan", "date"),
                           C("nama", "Nama", "ref", src="karyawan"),
                           C("pasangan", "Pasangan", "ref", src="karyawan"),
                           C("kode_motif", "Kode Motif", "ref", src="motif", req=1),
                           C("motif", "Motif", "auto", sql="r_kode_motif.nama_motif", hc="FF0000"),
                           C("jenis", "Jenis", "auto", sql="rj.nama", hc="FF0000"),
                           C("link_masuk", "Link Barang Masuk", "linkref", src="masuklink",
                             labelsql="""(SELECT d.sstb||' • '||mo.kode_motif||' ('||i.jumlah||') #'||i.id
                                 FROM dokumen_item i JOIN dokumen d ON d.id=i.dokumen_id
                                 JOIN ref_motif mo ON mo.id=i.motif_id WHERE i.id=t.link_masuk_id)""")]
                          + [C(c, c.capitalize(), "int") for c in CATS]
                          + [C("jumlah", "Jumlah", "auto", sql=JUMLAH_EXPR, num=1, hc="FF0000"),
                             C("ket", "Ket", "ref", src="ketprod"), C("catatan", "Catatan")]),
}

def dbcol(c):
    return c["n"] + "_id" if c["t"] in ("ref", "linkref") else c["n"]

def parse_link_id(v):
    """Ambil id dari label seperti '... #123' (dipakai kolom Link, opsional)."""
    m = re.search(r"#(\d+)\s*$", str(v or ""))
    return int(m.group(1)) if m else None

def dmy(v):
    try:
        return dt.date.fromisoformat(v).strftime("%d/%m/%Y")
    except Exception:
        return v or ""

def ensure_columns(db):
    """Tambah kolom baru (fitur Link opsional) ke tabel lama tanpa mengubah data yang sudah ada."""
    add = {"produksi": [("link_masuk_id", "INTEGER")],
           "dokumen_item": [("link_produksi_id", "INTEGER")]}
    for tbl, cols in add.items():
        have = {r[1] for r in db.execute(f"PRAGMA table_info({tbl})")}
        for name, decl in cols:
            if name not in have:
                db.execute(f"ALTER TABLE {tbl} ADD COLUMN {name} {decl}")

def init_db():
    os.makedirs(DATA_DIR, exist_ok=True)
    old = read_old()
    if old:
        os.makedirs(BACKUP_DIR, exist_ok=True)
        shutil.copy(DB_PATH, os.path.join(BACKUP_DIR, f"pra-migrasi-{dt.datetime.now():%Y%m%d-%H%M%S}.db"))
        os.remove(DB_PATH)
    db = sqlite3.connect(DB_PATH)
    db.execute("PRAGMA foreign_keys=ON")
    opening_migrated = False
    ocols = [r[1] for r in db.execute("PRAGMA table_info(opening)")]
    if ocols and "kode_motif_id" not in ocols:  # versi lama: Stok Awal per Jenis
        os.makedirs(BACKUP_DIR, exist_ok=True)
        bk = sqlite3.connect(os.path.join(BACKUP_DIR, f"pra-stokawal-{dt.datetime.now():%Y%m%d-%H%M%S}.db"))
        db.backup(bk); bk.close()
        db.execute("DROP INDEX IF EXISTS ux_opening")
        db.execute("ALTER TABLE opening RENAME TO opening_lama")
        opening_migrated = True
    for key, s in SPECS.items():
        cols = []
        for c in s["cols"]:
            if c["t"] == "auto":
                continue
            if c["t"] == "ref":
                cols.append(f"{c['n']}_id INTEGER REFERENCES {REFS[c['src']][0]}(id)")
            elif c["t"] == "linkref":
                cols.append(f"{c['n']}_id INTEGER")
            elif c["t"] == "int":
                cols.append(f"{c['n']} INTEGER NOT NULL DEFAULT 0")
            else:
                cols.append(f"{c['n']} TEXT")
        db.execute(f"CREATE TABLE IF NOT EXISTS {s['table']} (id INTEGER PRIMARY KEY AUTOINCREMENT, {', '.join(cols)})")
    db.execute("""CREATE TABLE IF NOT EXISTS dokumen (
        id INTEGER PRIMARY KEY AUTOINCREMENT, arah TEXT NOT NULL CHECK(arah IN ('M','K')),
        sstb TEXT NOT NULL, tanggal TEXT NOT NULL,
        dept_id INTEGER NOT NULL REFERENCES ref_dept(id),
        pengrajin_id INTEGER REFERENCES ref_pengrajin(id))""")
    db.execute("""CREATE TABLE IF NOT EXISTS dokumen_item (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        dokumen_id INTEGER NOT NULL REFERENCES dokumen(id) ON DELETE CASCADE,
        motif_id INTEGER NOT NULL REFERENCES ref_motif(id),
        ket_id INTEGER NOT NULL REFERENCES ref_ket(id),
        jumlah INTEGER NOT NULL CHECK(jumlah>0),
        link_produksi_id INTEGER)""")
    ensure_columns(db)  # tambah kolom baru ke tabel lama tanpa menghapus data
    db.execute("CREATE TABLE IF NOT EXISTS settings (k TEXT PRIMARY KEY, v TEXT)")
    db.execute("CREATE UNIQUE INDEX IF NOT EXISTS ux_dokumen ON dokumen(arah, sstb)")
    db.execute("CREATE INDEX IF NOT EXISTS ix_dokumen_tgl ON dokumen(tanggal)")
    db.execute("CREATE INDEX IF NOT EXISTS ix_item_dok ON dokumen_item(dokumen_id)")
    db.execute("CREATE INDEX IF NOT EXISTS ix_item_motif ON dokumen_item(motif_id)")
    db.execute("CREATE UNIQUE INDEX IF NOT EXISTS ux_opening_motif ON opening(kode_motif_id)")
    db.execute("CREATE INDEX IF NOT EXISTS ix_dokumen_arah_tgl ON dokumen(arah, tanggal)")
    db.execute("CREATE INDEX IF NOT EXISTS ix_prod_tgl ON produksi(tgl_produksi)")
    if old:
        import_old(db, old)
    else:
        seed(db)
    if opening_migrated:
        for jid, jml, tgl in db.execute("SELECT jenis_id, jumlah, tanggal_mulai FROM opening_lama").fetchall():
            opening_to_motif(db, jid, jml, tgl)
    db.commit()
    db.close()

def opening_to_motif(db, jenis_id, jumlah, tgl):
    """Stok awal lama per Jenis -> per Kode Motif. Otomatis hanya bila Jenis itu punya tepat 1 motif.
    Selain itu disimpan di tabel opening_lama agar tidak hilang; isi ulang manual di DB > Stok Awal Sistem."""
    ms = db.execute("SELECT id FROM ref_motif WHERE jenis_id=?", (jenis_id,)).fetchall()
    if len(ms) == 1:
        db.execute("INSERT OR IGNORE INTO opening(kode_motif_id,jumlah,tanggal_mulai) VALUES(?,?,?)",
                   (ms[0][0], jumlah, tgl))
        return True
    db.execute("CREATE TABLE IF NOT EXISTS opening_lama (jenis_id INTEGER, jumlah INTEGER, tanggal_mulai TEXT)")
    if not db.execute("SELECT 1 FROM opening_lama WHERE jenis_id=? AND jumlah=?", (jenis_id, jumlah)).fetchone():
        db.execute("INSERT INTO opening_lama VALUES(?,?,?)", (jenis_id, jumlah, tgl))
    nm = db.execute("SELECT nama FROM ref_jenis WHERE id=?", (jenis_id,)).fetchone()
    print(f"  [!] Stok awal jenis {nm[0] if nm else jenis_id} = {jumlah} belum bisa dipindah otomatis "
          f"({len(ms)} motif). Isi ulang per Kode Motif di halaman DB > Stok Awal Sistem.")
    return False

def seed(db):
    if db.execute("SELECT COUNT(*) FROM ref_ket").fetchone()[0]:
        return
    ket = [("Wadimor", "WD"), ("Washing", "WS"), ("Celup", "CLP"), ("Sarung Anak", "SA"),
           ("Menclek", "MC"), ("Rusak", "RSK"), ("Belum Neci", "BNC"), ("Sudah Neci", "SNC"),
           ("Jahit 2 Belum Jahit", "J2B"), ("Jahit 2 Sudah Jahit", "J2S"),
           ("Wadimor Sudah Jahit", "WSJ"), ("Junior Sudah Jahit", "JSJ"),
           ("D Biasa Belum Jahit", "DBB"), ("D Biasa Sudah Jahit", "DBS"),
           ("Belum Sensor Grade Harian", "BSGH"), ("Belum Sensor Kelolokan", "BSK"),
           ("Belum Sensor Tindesan", "BST"), ("Belum Sensor Neci", "BSN"),
           ("Belum Sensor Perbaikan", "BSP"), ("Lolos Jahit", "LJ"), ("Cabut", "CBT"),
           ("Perbaikan", "MC")]  # NB: kode MC dipakai ganda di file asli (Menclek & Perbaikan)
    db.executemany("INSERT INTO ref_ket(nama,kode) VALUES (?,?)", ket)
    db.executemany("INSERT INTO ref_dept(nama,kode) VALUES (?,?)", [
        ("PEMOTONGAN", "PMTG"), ("JAHIT PACK1", "JAHIT1"), ("WIP2 PACK1", "WIP2 1"),
        ("JAHIT PACK2", "JAHIT2"), ("GRADE", "GRADE"), ("JASA LUAR", "JASA LUAR"),
        ("FINISHING", "FINISHING"), ("LAIN-LAIN", "LAIN-LAIN"), ("JAHIT LUAR", "JAHIT LUAR")])
    db.executemany("INSERT INTO ref_ketprod(nama) VALUES (?)", [(x,) for x in [
        "WEAVING", "DYING", "FINISHING", "NECI", "MENCLEK", "SENSOR ULANG", "D LAIN-LAIN", "AFTER WASHING"]])
    db.executemany("INSERT INTO ref_pengrajin(nama) VALUES (?)", [(x,) for x in [
        "SALEH BA'AGIL", "SUHARNO", "DYAH AYU SHINTA", "JUNTORO AJI", "MULTAZAM", "CASWITO", "AISAH"]])
    db.executemany("INSERT INTO ref_karyawan(nik,nama) VALUES (?,?)", [
        ("6364350614", "JUMIATI"), ("9288651120", "SARJINEM"), ("11427471022", "LAILATUL HIKMAH"),
        ("5764350214", "LISAH KURNIAWATI"), ("11512291122", "EKKA PUSPITA SARI"),
        ("13462470824", "NUR AZIZAH"), ("0847570805", "NUR HIDAYAH"),
        ("13238370324", "SITI FADHILAH"), ("13445470824", "DEVI INDRIYANI"),
        ("10306290221", "FIKA ALFIANTIN")])
    db.executemany("INSERT INTO ref_jenis(nama) VALUES (?)", [("30 STR",), ("SULAM",)])
    j = db.execute("SELECT id FROM ref_jenis WHERE nama='30 STR'").fetchone()[0]
    db.executemany("INSERT INTO ref_motif(kode_motif,nama_motif,jenis_id) VALUES (?,?,?)", [
        ("0000.01", "30 STR (GRADE D)", j), ("0021.01", "Horison Gradasi", j),
        ("0123.01", "Yunior Tanggung (30 STR)", j), ("0061.01", "Horison", j), ("0045.01", "Tiga Dara", j)])

# ---- migrasi dari versi lama (data teks ganda -> ID master) --------------
OLD_TABLES = ("ref_ket", "ref_dept", "ref_ketprod", "ref_pengrajin", "ref_karyawan",
              "ref_motif", "opening", "produksi", "trx_masuk", "trx_keluar")

def read_old():
    if not os.path.exists(DB_PATH):
        return None
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    try:
        cols = [r["name"] for r in con.execute("PRAGMA table_info(ref_motif)")]
        if not cols or "jenis_id" in cols:
            return None
        data = {}
        for t in OLD_TABLES:
            try:
                data[t] = [dict(r) for r in con.execute(f"SELECT * FROM {t} ORDER BY id")]
            except sqlite3.OperationalError:
                data[t] = []
        try:
            r = con.execute("SELECT v FROM settings WHERE k='strict_keluar'").fetchone()
            data["strict"] = r["v"] if r else None
        except sqlite3.OperationalError:
            data["strict"] = None
        return data
    finally:
        con.close()

def import_old(db, o):
    def ins(sql, args):
        return db.execute(sql, args).lastrowid
    m = {k: {} for k in ("ket", "dept", "ketprod", "pengrajin", "karyawan", "jenis", "motif")}
    for r in o["ref_ket"]:
        m["ket"][r["nama"].lower()] = ins("INSERT INTO ref_ket(nama,kode) VALUES(?,?)", (r["nama"], r["kode"]))
    for r in o["ref_dept"]:
        m["dept"][r["nama"].lower()] = ins("INSERT INTO ref_dept(nama,kode) VALUES(?,?)", (r["nama"], r["kode"]))
    for r in o["ref_ketprod"]:
        m["ketprod"][r["nama"].lower()] = ins("INSERT INTO ref_ketprod(nama) VALUES(?)", (r["nama"],))
    for r in o["ref_pengrajin"]:
        m["pengrajin"][r["nama"].lower()] = ins("INSERT INTO ref_pengrajin(nama) VALUES(?)", (r["nama"],))
    for r in o["ref_karyawan"]:
        m["karyawan"][r["nama"].lower()] = ins("INSERT INTO ref_karyawan(nik,nama) VALUES(?,?)", (r["nik"], r["nama"]))
    for j in sorted({(r.get("jenis") or "").upper() for r in o["ref_motif"] + o["opening"]} - {""}):
        m["jenis"][j.lower()] = ins("INSERT INTO ref_jenis(nama) VALUES(?)", (j,))
    for r in o["ref_motif"]:
        m["motif"][r["kode_motif"].lower()] = ins(
            "INSERT INTO ref_motif(kode_motif,nama_motif,jenis_id) VALUES(?,?,?)",
            (r["kode_motif"], r["nama_motif"], m["jenis"].get((r["jenis"] or "").lower())))
    for r in o["opening"]:
        jid = m["jenis"].get((r["jenis"] or "").lower())
        if jid:
            opening_to_motif(db, jid, r["jumlah"], r["tanggal_mulai"])
    for r in o["produksi"]:
        pas = None
        if r.get("pasangan"):
            k = r["pasangan"].strip().lower()
            pas = m["karyawan"].get(k) or ins("INSERT INTO ref_karyawan(nama) VALUES(?)", (r["pasangan"].strip(),))
            m["karyawan"][k] = pas
        mid = m["motif"].get((r["kode_motif"] or "").lower())
        if not mid:
            continue
        cols = ["tgl_produksi", "tgl_masuk_sanggan", "nama_id", "pasangan_id", "kode_motif_id"] + CATS + ["ket_id", "catatan"]
        vals = [r["tgl_produksi"], r["tgl_masuk_sanggan"], m["karyawan"].get((r["nama"] or "").lower()), pas, mid] \
               + [r.get(c) or 0 for c in CATS] + [m["ketprod"].get((r["ket"] or "").lower()), r["catatan"]]
        db.execute(f"INSERT INTO produksi({','.join(cols)}) VALUES({','.join('?'*len(cols))})", vals)
    for arah, tbl in (("M", "trx_masuk"), ("K", "trx_keluar")):
        docs, used = {}, set()
        for r in o[tbl]:
            dept = m["dept"].get((r["dept"] or "").lower())
            mid = m["motif"].get((r["kode_motif"] or "").lower())
            kid = m["ket"].get((r["ket"] or "").lower())
            if not (dept and mid and kid):
                continue
            peng = m["pengrajin"].get((r["pengrajin"] or "").lower())
            hk = (r["sstb"], r["tanggal"], dept, peng)
            did = docs.get(hk)
            if did is None:  # SSTB sama tapi tanggal/dept beda pada data lama -> beri akhiran agar tetap unik
                name, n = r["sstb"], 1
                while name.lower() in used:
                    n += 1
                    name = f"{r['sstb']} ({n})"
                used.add(name.lower())
                did = docs[hk] = ins("INSERT INTO dokumen(arah,sstb,tanggal,dept_id,pengrajin_id) VALUES(?,?,?,?,?)",
                                     (arah, name, r["tanggal"], dept, peng))
            db.execute("INSERT INTO dokumen_item(dokumen_id,motif_id,ket_id,jumlah) VALUES(?,?,?,?)",
                       (did, mid, kid, r["jumlah"]))
    if o.get("strict"):
        db.execute("INSERT OR REPLACE INTO settings(k,v) VALUES('strict_keluar',?)", (o["strict"],))

# --------------------------------------------------------------------------
# Util
# --------------------------------------------------------------------------
def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys=ON")
    return g.db

@app.teardown_appcontext
def close_db(_):
    d = g.pop("db", None)
    if d:
        d.close()

def parse_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    if re.match(r"^\d{4}-\d{2}-\d{2}", s):
        s = s[:10]
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%b-%Y", "%d %b %Y"):
        try:
            return dt.datetime.strptime(s, f).date().isoformat()
        except ValueError:
            pass
    return None

def parse_int(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else None
    s = str(v).strip().replace(" ", "")
    if re.match(r"^\d{1,3}(\.\d{3})+$", s):
        s = s.replace(".", "")
    try:
        f = float(s)
    except ValueError:
        return None
    return int(f) if f.is_integer() else None

def text_of(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()

def get_setting(db, k, d=""):
    r = db.execute("SELECT v FROM settings WHERE k=?", (k,)).fetchone()
    return r["v"] if r else d

def ref_lookup(db, src, val):
    tbl, col = REFS[src]
    return db.execute(f"SELECT * FROM {tbl} WHERE lower({col})=lower(?)", (str(val),)).fetchone()

PAGE_DEFAULT, PAGE_MAX = 25, 100   # baris per halaman (agar tabel tidak berat)
DEFAULT_DAYS = 7                   # rentang tanggal bawaan bila tidak dikirim

def default_from():
    return (dt.date.today() - dt.timedelta(days=DEFAULT_DAYS - 1)).isoformat()

def page_args(args):
    try:
        size = int(args.get("size", PAGE_DEFAULT))
    except ValueError:
        size = PAGE_DEFAULT
    try:
        page = int(args.get("page", 1))
    except ValueError:
        page = 1
    return max(1, min(size, PAGE_MAX)), max(page, 1)

def fmt_id(n):
    return f"{n:,}".replace(",", ".")

# --------------------------------------------------------------------------
# CRUD generik (master + produksi)
# --------------------------------------------------------------------------
def spec_or_404(key):
    if key not in SPECS:
        abort(404)
    return SPECS[key]

def select_parts(key):
    s = SPECS[key]
    sel, joins = ["t.id AS id"], []
    for c in s["cols"]:
        n = c["n"]
        if c["t"] == "ref":
            tbl, col = REFS[c["src"]]
            joins.append(f"LEFT JOIN {tbl} r_{n} ON r_{n}.id=t.{n}_id")
            sel.append(f"r_{n}.{col} AS {n}")
        elif c["t"] == "auto":
            sel.append(f"({c['sql']}) AS {n}")
        elif c["t"] == "linkref":
            sel.append(f"({c['labelsql']}) AS {n}")
        else:
            sel.append(f"t.{n} AS {n}")
    joins += s.get("joins", [])
    return ", ".join(sel), " ".join(joins)

def col_info(key):
    """nama kolom -> (ekspresi SQL, tipe). Dipakai untuk pencarian (q), filter per-kolom ala Excel, dan sort."""
    out = {}
    for c in SPECS[key]["cols"]:
        n = c["n"]
        if c["t"] == "ref":
            tbl, col = REFS[c["src"]]
            out[n] = (f"r_{n}.{col}", c["t"])
        elif c["t"] == "auto":
            out[n] = (f"({c['sql']})", c["t"])
        elif c["t"] == "linkref":
            out[n] = (f"({c['labelsql']})", c["t"])
        else:
            out[n] = (f"t.{n}", c["t"])
    return out

def search_exprs(key):
    """Ekspresi SQL kolom teks, dipakai untuk kotak pencarian (q). Kolom angka dilewati."""
    return [e for e, t in col_info(key).values() if t != "int"]

def parse_filters(args):
    try:
        return json.loads(args.get("filters") or "{}") or {}
    except (TypeError, ValueError):
        return {}

def apply_col_filters(where, params, info, flt):
    """Filter per-kolom ala Excel: {nama_kolom: [nilai, ...]} -> WHERE ekspresi IN (...)."""
    for col, vals in (flt or {}).items():
        if col in info and vals:
            expr, _ = info[col]
            where.append(f"{expr} IN ({','.join('?' * len(vals))})")
            params.extend(vals)

def parse_ranges(args):
    try:
        return json.loads(args.get("ranges") or "{}") or {}
    except (TypeError, ValueError):
        return {}

def apply_ranges(where, params, info, ranges):
    """Filter rentang angka: {nama_kolom: {min, max}} -> WHERE ekspresi BETWEEN."""
    for col, rng in (ranges or {}).items():
        if col not in info or not isinstance(rng, dict):
            continue
        expr, _ = info[col]
        mn, mx = rng.get("min"), rng.get("max")
        if mn not in (None, ""):
            where.append(f"{expr}>=?"); params.append(mn)
        if mx not in (None, ""):
            where.append(f"{expr}<=?"); params.append(mx)

def group_order(args, names, order):
    """Prefiks ORDER BY dengan kolom Group By (ASC) supaya baris sekelompok berurutan di halaman yang sama."""
    g = args.get("group")
    if g and g in names:
        return f"{g} ASC, {order}"
    return order

JENIS_FILTER = {  # key -> ekspresi SQL id jenis milik baris, untuk filter "kategori"
    "motif": "t.jenis_id",
    "opening": "(SELECT jenis_id FROM ref_motif WHERE id=t.kode_motif_id)",
    "produksi": "(SELECT jenis_id FROM ref_motif WHERE id=t.kode_motif_id)",
}

def list_filters(key, args):
    """joins, daftar klausa WHERE, params — dipakai bareng oleh list_query dan /api/distinct."""
    s = SPECS[key]
    sel, joins = select_parts(key)
    where, params = [], []
    dc = s.get("datecol")
    if dc:
        where.append(f"t.{dc}>=?"); params.append(args.get("from") or default_from())
        if args.get("to"):
            where.append(f"t.{dc}<=?"); params.append(args["to"])
    if args.get("jenis") and key in JENIS_FILTER:
        where.append(f"{JENIS_FILTER[key]}=(SELECT id FROM ref_jenis WHERE nama=?)")
        params.append(args["jenis"])
    q = (args.get("q") or "").strip()
    if q:
        exprs = search_exprs(key)
        if exprs:
            where.append("(" + " OR ".join(f"{e} LIKE ?" for e in exprs) + ")")
            params += [f"%{q}%"] * len(exprs)
    info = col_info(key)
    apply_col_filters(where, params, info, parse_filters(args))
    apply_ranges(where, params, info, parse_ranges(args))
    return sel, joins, where, params

def list_query(key, args):
    s = SPECS[key]
    sel, joins, where, params = list_filters(key, args)
    w = ("WHERE " + " AND ".join(where)) if where else ""
    names = [c["n"] for c in s["cols"]]
    dc = s.get("datecol")
    direction = "ASC" if args.get("dir") == "asc" else "DESC"
    if args.get("sort") in names:
        order = f"{args['sort']} {direction}, t.id DESC"
    elif dc:
        order = f"t.{dc} DESC, t.id DESC"
    else:
        order = "t.id ASC"
    order = group_order(args, names, order)
    return sel, joins, w, params, order

@app.get("/api/distinct/<key>")
def api_distinct(key):
    spec_or_404(key)
    info = col_info(key)
    col = request.args.get("col", "")
    if col not in info:
        return jsonify([])
    args = request.args.to_dict()
    flt = parse_filters(args)
    flt.pop(col, None)  # kolom yang sedang dibuka tidak membatasi daftar pilihannya sendiri
    args["filters"] = json.dumps(flt)
    _, joins, where, params = list_filters(key, args)
    expr, _ = info[col]
    where = where + [f"{expr} IS NOT NULL", f"{expr}<>''"]
    w = "WHERE " + " AND ".join(where)
    rows = get_db().execute(
        f"SELECT DISTINCT {expr} AS v FROM {SPECS[key]['table']} t {joins} {w} ORDER BY {expr} LIMIT 500",
        params).fetchall()
    return jsonify([r["v"] for r in rows])

@app.get("/api/meta/<key>")
def api_meta(key):
    s = spec_or_404(key)
    return jsonify(key=key, title=s["title"], datecol=s.get("datecol"), cats=CATS,
                   cols=[{k: v for k, v in c.items() if k in ("n", "l", "t", "src", "req", "num")} for c in s["cols"]])

OPT_LIMIT_DEFAULT, OPT_LIMIT_MAX = 50, 200  # rekomendasi input: tidak pernah kirim semua data sekaligus

def search_masuklink(db, q, limit):
    where, params = ["d.arah='M'"], []
    fid = parse_link_id(q) if q else None
    if fid:  # q berupa label lengkap (mis. saat validasi link yang sudah tersimpan) -> cari persis id-nya
        where.append("i.id=?"); params.append(fid)
    elif q:
        where.append("(d.sstb LIKE ? OR mo.kode_motif LIKE ? OR mo.nama_motif LIKE ?)")
        params += [f"%{q}%"] * 3
    rows = db.execute(f"""SELECT i.id, d.sstb, d.tanggal, mo.kode_motif, mo.nama_motif, i.jumlah
        FROM dokumen_item i JOIN dokumen d ON d.id=i.dokumen_id JOIN ref_motif mo ON mo.id=i.motif_id
        WHERE {' AND '.join(where)} ORDER BY d.tanggal DESC, i.id DESC LIMIT ?""", params + [limit]).fetchall()
    return [dict(v=f"{r['sstb']} • {r['kode_motif']} ({fmt_id(r['jumlah'])}) #{r['id']}",
                sub=f"{r['nama_motif']} · {dmy(r['tanggal'])}") for r in rows]

def search_produksilink(db, q, limit):
    where, params = [], []
    fid = parse_link_id(q) if q else None
    if fid:  # q berupa label lengkap -> cari persis id-nya (validasi link yang sudah tersimpan)
        where.append("p.id=?"); params.append(fid)
    elif q:
        where.append("(pm.kode_motif LIKE ? OR pm.nama_motif LIKE ? OR kr.nama LIKE ?)")
        params += [f"%{q}%"] * 3
    w = ("WHERE " + " AND ".join(where)) if where else ""
    rows = db.execute(f"""SELECT p.id, p.tgl_produksi, pm.kode_motif, pm.nama_motif, kr.nama nama,
            ({JUMLAH_EXPR.replace('t.', 'p.')}) jumlah
        FROM produksi p LEFT JOIN ref_motif pm ON pm.id=p.kode_motif_id
        LEFT JOIN ref_karyawan kr ON kr.id=p.nama_id
        {w} ORDER BY p.tgl_produksi DESC, p.id DESC LIMIT ?""", params + [limit]).fetchall()
    return [dict(v=f"{r['kode_motif'] or '-'} • {dmy(r['tgl_produksi'])} ({fmt_id(r['jumlah'] or 0)}) #{r['id']}",
                sub=f"{r['nama_motif'] or ''} · {r['nama'] or ''}") for r in rows]

@app.get("/api/opt/<src>")
def api_opt(src):
    db = get_db()
    q = (request.args.get("q") or "").strip()
    try:
        limit = max(1, min(int(request.args.get("limit", OPT_LIMIT_DEFAULT)), OPT_LIMIT_MAX))
    except ValueError:
        limit = OPT_LIMIT_DEFAULT
    if src == "masuklink":
        return jsonify(search_masuklink(db, q, limit))
    if src == "produksilink":
        return jsonify(search_produksilink(db, q, limit))
    if src not in REFS:
        abort(404)
    out = []
    if src == "motif":
        where, params = "", []
        if q:
            where = "WHERE m.kode_motif LIKE ? OR m.nama_motif LIKE ? OR j.nama LIKE ?"
            params = [f"%{q}%"] * 3
        for r in db.execute(f"""SELECT m.kode_motif, m.nama_motif, j.nama jenis FROM ref_motif m
                               LEFT JOIN ref_jenis j ON j.id=m.jenis_id {where}
                               ORDER BY m.kode_motif LIMIT ?""", params + [limit]):
            out.append(dict(v=r["kode_motif"], sub=f"{r['nama_motif']} · {r['jenis'] or ''}",
                            motif=r["nama_motif"], jenis=r["jenis"]))
    else:
        tbl, col = REFS[src]
        extra = ["nik"] if src == "karyawan" else ["kode"] if src == "ket" else []
        where, params = "", []
        if q:
            cs = [col] + extra
            where = "WHERE " + " OR ".join(f"{c} LIKE ?" for c in cs)
            params = [f"%{q}%"] * len(cs)
        for r in db.execute(f"SELECT * FROM {tbl} {where} ORDER BY {col} LIMIT ?", params + [limit]):
            d = dict(v=r[col])
            if src == "ket":
                d.update(sub=r["kode"], kode=r["kode"])
            elif src == "karyawan":
                d.update(sub=r["nik"] or "")
            out.append(d)
    return jsonify(out)

@app.get("/api/<key>")
def api_list(key):
    s = spec_or_404(key)
    db = get_db()
    sel, joins, w, params, order = list_query(key, request.args)
    size, page = page_args(request.args)
    tsum = None
    if key == "produksi":
        total, tsum = db.execute(
            f"SELECT COUNT(*), COALESCE(SUM({JUMLAH_EXPR}),0) FROM produksi t {joins} {w}", params).fetchone()
    else:
        total = db.execute(f"SELECT COUNT(*) FROM {s['table']} t {joins} {w}", params).fetchone()[0]
    rows = db.execute(f"SELECT {sel} FROM {s['table']} t {joins} {w} ORDER BY {order} LIMIT ? OFFSET ?",
                      params + [size, (page - 1) * size]).fetchall()
    return jsonify(rows=[dict(r) for r in rows], total=total, pages=max(1, -(-total // size)), sum=tsum)

def clean_row(db, key, row, rid=None):
    spec = SPECS[key]
    out, errs = {}, []
    for c in spec["cols"]:
        n, t = c["n"], c["t"]
        if t == "auto":
            continue
        v = row.get(n)
        if isinstance(v, str):
            v = v.strip()
        if v in ("", None):
            v = None
        bad = False
        if t == "date" and v is not None:
            pv = parse_date(v)
            if pv is None:
                errs.append(f"{c['l']}: tanggal tidak valid"); bad = True
            v = pv
        elif t == "int":
            if v is None:
                v = None if c.get("req") else 0
            else:
                pv = parse_int(v)
                if pv is None or pv < 0:
                    errs.append(f"{c['l']}: harus bilangan bulat >= 0"); bad = True; v = None
                else:
                    v = pv
        elif t == "ref" and v is not None:
            r = ref_lookup(db, c["src"], v)
            if r is None:
                errs.append(f"{c['l']}: '{v}' tidak ada di master (DB)"); bad = True; v = None
            else:
                v = r["id"]
        elif t == "linkref":
            v = parse_link_id(v) if v is not None else None  # opsional: teks tak dikenal = tidak ditautkan
        elif v is not None:
            v = text_of(v)
            if c.get("upper"):
                v = v.upper()
        if c.get("req") and v is None and not bad:
            errs.append(f"{c['l']} wajib diisi")
        out[dbcol(c)] = v
    for c in spec["cols"]:  # keunikan
        col = dbcol(c)
        if c.get("uniq") and out.get(col) is not None:
            if c["t"] == "ref":
                q, a = f"{col}=?", out[col]
            else:
                q, a = f"lower({col})=lower(?)", str(out[col])
            if db.execute(f"SELECT id FROM {spec['table']} WHERE {q} AND id<>?", (a, rid or -1)).fetchone():
                errs.append(f"{c['l']}: '{row.get(c['n'])}' sudah ada")
    return out, errs

def save(key, rid=None):
    s = spec_or_404(key)
    db = get_db()
    out, errs = clean_row(db, key, request.get_json(force=True) or {}, rid)
    if errs:
        return jsonify(ok=False, errors=errs), 400
    cols = list(out.keys())
    if rid is None:
        db.execute(f"INSERT INTO {s['table']}({','.join(cols)}) VALUES ({','.join('?'*len(cols))})",
                   [out[c] for c in cols])
    else:
        db.execute(f"UPDATE {s['table']} SET {','.join(c + '=?' for c in cols)} WHERE id=?",
                   [out[c] for c in cols] + [rid])
    db.commit()
    return jsonify(ok=True, warnings=[])

@app.post("/api/<key>")
def api_create(key):
    return save(key)

@app.put("/api/<key>/<int:rid>")
def api_update(key, rid):
    return save(key, rid)

@app.delete("/api/<key>/<int:rid>")
def api_delete(key, rid):
    s = spec_or_404(key)
    db = get_db()
    used = 0
    for tbl, col in USAGE.get(key, []):
        used += db.execute(f"SELECT COUNT(*) FROM {tbl} WHERE {col}=?", (rid,)).fetchone()[0]
    if used:
        return jsonify(ok=False, errors=[f"Tidak bisa dihapus: masih dipakai di {used} data. Ubah namanya saja bila perlu."]), 409
    db.execute(f"DELETE FROM {s['table']} WHERE id=?", (rid,))
    db.commit()
    return jsonify(ok=True)

@app.route("/api/setting/strict", methods=["GET", "POST"])
def api_strict():
    db = get_db()
    if request.method == "POST":
        v = "1" if (request.get_json(force=True) or {}).get("on") else "0"
        db.execute("INSERT OR REPLACE INTO settings(k,v) VALUES('strict_keluar',?)", (v,))
        db.commit()
    return jsonify(on=get_setting(db, "strict_keluar") == "1")

# --------------------------------------------------------------------------
# Dokumen masuk / keluar (satu SSTB = satu dokumen, banyak baris barang)
# --------------------------------------------------------------------------
ARAH = {"masuk": "M", "keluar": "K"}
DOC_TITLE = {"masuk": "Barang Masuk", "keluar": "Barang Keluar"}
DOC_COLOR = {"masuk": "ED7D31", "keluar": "70AD47"}

LINK_PRODUKSI_LABEL_SQL = """(SELECT pm.kode_motif||' • '||p.tgl_produksi||' #'||p.id
    FROM produksi p LEFT JOIN ref_motif pm ON pm.id=p.kode_motif_id WHERE p.id=i.link_produksi_id)"""
LINE_SEL = f"""d.sstb AS sstb, d.tanggal AS tanggal, dp.nama AS dept, m.kode_motif AS kode_motif,
    m.nama_motif AS motif, j.nama AS jenis, i.jumlah AS jumlah, k.nama AS ket,
    pg.nama AS pengrajin, k.kode AS rumus, {LINK_PRODUKSI_LABEL_SQL} AS link_produksi,
    i.id AS id, d.id AS doc_id"""
LINE_FROM = """FROM dokumen_item i JOIN dokumen d ON d.id=i.dokumen_id
    JOIN ref_dept dp ON dp.id=d.dept_id JOIN ref_motif m ON m.id=i.motif_id
    LEFT JOIN ref_jenis j ON j.id=m.jenis_id JOIN ref_ket k ON k.id=i.ket_id
    LEFT JOIN ref_pengrajin pg ON pg.id=d.pengrajin_id"""
LINE_COLS = [("sstb", "SSTB", ""), ("tanggal", "Tanggal", "date"), ("dept", "Dept", ""),
             ("kode_motif", "Kode Motif", ""), ("motif", "Motif", ""), ("jenis", "Jenis", ""),
             ("jumlah", "Jumlah", "int"), ("ket", "Ket", ""), ("pengrajin", "Nama Pengrajin", ""),
             ("rumus", "Rumus", ""), ("link_produksi", "Link Produksi", "")]

def arah_or_404(name):
    if name not in ARAH:
        abort(404)
    return ARAH[name]

LINE_SEARCH_EXPRS = ["d.sstb", "dp.nama", "m.kode_motif", "m.nama_motif", "j.nama", "k.nama", "pg.nama", "k.kode"]
LINE_COL_EXPRS = {"sstb": "d.sstb", "tanggal": "d.tanggal", "dept": "dp.nama", "kode_motif": "m.kode_motif",
                   "motif": "m.nama_motif", "jenis": "j.nama", "jumlah": "i.jumlah", "ket": "k.nama",
                   "pengrajin": "pg.nama", "rumus": "k.kode", "link_produksi": LINK_PRODUKSI_LABEL_SQL}

def lines_filters(arah, args):
    """daftar klausa WHERE, params — dipakai bareng oleh lines_query dan /api/distinct/lines."""
    where, params = ["d.arah=?"], [arah]
    where.append("d.tanggal>=?"); params.append(args.get("from") or default_from())
    if args.get("to"):
        where.append("d.tanggal<=?"); params.append(args["to"])
    if args.get("dept"):
        where.append("dp.nama=?"); params.append(args["dept"])
    q = (args.get("q") or "").strip()
    if q:
        where.append("(" + " OR ".join(f"{e} LIKE ?" for e in LINE_SEARCH_EXPRS) + ")")
        params += [f"%{q}%"] * len(LINE_SEARCH_EXPRS)
    line_info = {k: (v, "int" if k == "jumlah" else "") for k, v in LINE_COL_EXPRS.items()}
    apply_col_filters(where, params, line_info, parse_filters(args))
    apply_ranges(where, params, line_info, parse_ranges(args))
    return where, params

def lines_query(arah, args):
    where, params = lines_filters(arah, args)
    w = "WHERE " + " AND ".join(where)
    names = [c[0] for c in LINE_COLS]
    direction = "ASC" if args.get("dir") == "asc" else "DESC"
    if args.get("sort") in names:
        order = f"{args['sort']} {direction}, doc_id DESC, id ASC"
    else:
        order = "tanggal DESC, doc_id DESC, id ASC"
    order = group_order(args, names, order)
    return w, params, order

@app.get("/api/distinct/lines/<name>")
def api_distinct_lines(name):
    arah = arah_or_404(name)
    col = request.args.get("col", "")
    if col not in LINE_COL_EXPRS:
        return jsonify([])
    args = request.args.to_dict()
    flt = parse_filters(args)
    flt.pop(col, None)
    args["filters"] = json.dumps(flt)
    where, params = lines_filters(arah, args)
    expr = LINE_COL_EXPRS[col]
    where = where + [f"{expr} IS NOT NULL", f"{expr}<>''"]
    w = "WHERE " + " AND ".join(where)
    rows = get_db().execute(f"SELECT DISTINCT {expr} AS v {LINE_FROM} {w} ORDER BY {expr} LIMIT 500", params).fetchall()
    return jsonify([r["v"] for r in rows])

@app.get("/api/lines/<name>")
def api_lines(name):
    arah = arah_or_404(name)
    db = get_db()
    w, params, order = lines_query(arah, request.args)
    size, page = page_args(request.args)
    total, tsum, docs = db.execute(
        f"SELECT COUNT(*), COALESCE(SUM(i.jumlah),0), COUNT(DISTINCT d.id) {LINE_FROM} {w}", params).fetchone()
    rows = db.execute(f"SELECT {LINE_SEL} {LINE_FROM} {w} ORDER BY {order} LIMIT ? OFFSET ?",
                      params + [size, (page - 1) * size]).fetchall()
    return jsonify(rows=[dict(r) for r in rows], total=total, pages=max(1, -(-total // size)),
                   sum=tsum, docs=docs)

@app.get("/api/doc/<name>/<int:doc_id>")
def api_doc_get(name, doc_id):
    arah = arah_or_404(name)
    db = get_db()
    d = db.execute("""SELECT d.id, d.sstb, d.tanggal, dp.nama dept, pg.nama pengrajin FROM dokumen d
        JOIN ref_dept dp ON dp.id=d.dept_id LEFT JOIN ref_pengrajin pg ON pg.id=d.pengrajin_id
        WHERE d.id=? AND d.arah=?""", (doc_id, arah)).fetchone()
    if not d:
        abort(404)
    items = db.execute(f"""SELECT m.kode_motif, k.nama ket, i.jumlah,
            {LINK_PRODUKSI_LABEL_SQL} AS link_produksi
        FROM dokumen_item i JOIN ref_motif m ON m.id=i.motif_id JOIN ref_ket k ON k.id=i.ket_id
        WHERE i.dokumen_id=? ORDER BY i.id""", (doc_id,)).fetchall()
    out = dict(d)
    out["items"] = [dict(x) for x in items]
    return jsonify(out)

def balance_upto(db, jenis_id, tanggal, excl_doc=None):
    """Saldo per Jenis = jumlah stok awal semua motif pada jenis itu + masuk - keluar (mulai tanggal awal tiap motif)."""
    base = db.execute("""SELECT COALESCE(SUM(o.jumlah),0) FROM opening o JOIN ref_motif m ON m.id=o.kode_motif_id
                         WHERE m.jenis_id=?""", (jenis_id,)).fetchone()[0]

    def s(arah):
        return db.execute("""SELECT COALESCE(SUM(i.jumlah),0) FROM dokumen_item i
            JOIN dokumen d ON d.id=i.dokumen_id JOIN ref_motif m ON m.id=i.motif_id
            LEFT JOIN opening o ON o.kode_motif_id=m.id
            WHERE d.arah=? AND m.jenis_id=? AND d.tanggal>=COALESCE(o.tanggal_mulai,'') AND d.tanggal<=? AND d.id<>?""",
                          (arah, jenis_id, tanggal, excl_doc or -1)).fetchone()[0]
    return base + s("M") - s("K")

@app.get("/api/available")
def api_available():
    db = get_db()
    r = ref_lookup(db, "jenis", request.args.get("jenis", ""))
    tgl = parse_date(request.args.get("tanggal")) or dt.date.today().isoformat()
    doc = int(request.args.get("doc") or 0) or None
    return jsonify(available=balance_upto(db, r["id"], tgl, doc) if r else None)

def validate_doc(db, arah, data, doc_id=None):
    """Kembalikan (bersih, errors[(idx|None,msg)], warnings)."""
    errs, warns = [], []
    sstb = text_of(data.get("sstb"))
    if not sstb:
        errs.append((None, "SSTB wajib diisi"))
    elif db.execute("SELECT 1 FROM dokumen WHERE arah=? AND lower(sstb)=lower(?) AND id<>?",
                    (arah, sstb, doc_id or -1)).fetchone():
        errs.append((None, f"SSTB '{sstb}' sudah ada. Buka dari daftar untuk mengeditnya."))
    tanggal = parse_date(data.get("tanggal"))
    if not tanggal:
        errs.append((None, "Tanggal wajib diisi (format valid)"))
    dept = None
    if not text_of(data.get("dept")):
        errs.append((None, "Dept wajib diisi"))
    else:
        dept = ref_lookup(db, "dept", text_of(data["dept"]))
        if not dept:
            errs.append((None, f"Dept '{data['dept']}' tidak ada di master (DB)"))
    pg = None
    if text_of(data.get("pengrajin")):
        pg = ref_lookup(db, "pengrajin", text_of(data["pengrajin"]))
        if not pg:
            errs.append((None, f"Pengrajin '{data['pengrajin']}' tidak ada di master (DB)"))
    items, per_jenis = [], {}
    raw = data.get("items") or []
    if not raw:
        errs.append((None, "Minimal satu baris barang"))
    for i, it in enumerate(raw):
        m = ref_lookup(db, "motif", text_of(it.get("kode_motif"))) if text_of(it.get("kode_motif")) else None
        if not m:
            errs.append((i, f"Kode Motif '{text_of(it.get('kode_motif'))}' tidak ada di master (DB)"))
        k = ref_lookup(db, "ket", text_of(it.get("ket"))) if text_of(it.get("ket")) else None
        if not k:
            errs.append((i, f"Ket '{text_of(it.get('ket'))}' tidak ada di master (DB)"))
        jml = parse_int(it.get("jumlah")) if text_of(it.get("jumlah")) else None
        if jml is None or jml <= 0:
            errs.append((i, "Jumlah harus bilangan bulat > 0"))
        if m and k and jml and jml > 0:
            lp = parse_link_id(it.get("link_produksi")) if arah == "K" else None  # opsional
            items.append(dict(motif_id=m["id"], ket_id=k["id"], jumlah=jml, link_produksi_id=lp))
            per_jenis[m["jenis_id"]] = per_jenis.get(m["jenis_id"], 0) + jml
    if arah == "K" and not errs and tanggal:
        strict = get_setting(db, "strict_keluar") == "1"
        for jid, need in per_jenis.items():
            avail = balance_upto(db, jid, tanggal, doc_id)
            if need > avail:
                nm = db.execute("SELECT nama FROM ref_jenis WHERE id=?", (jid,)).fetchone()["nama"]
                msg = f"Stok {nm} per {tanggal} hanya {fmt_id(avail)}, sedangkan keluar {fmt_id(need)}"
                if strict:
                    errs.append((None, msg + " (mode ketat: ditolak)"))
                else:
                    warns.append(msg)
    clean = dict(sstb=sstb, tanggal=tanggal, dept_id=dept["id"] if dept else None,
                 pengrajin_id=pg["id"] if pg else None, items=items)
    return clean, errs, warns

def insert_doc(db, arah, c):
    did = db.execute("INSERT INTO dokumen(arah,sstb,tanggal,dept_id,pengrajin_id) VALUES(?,?,?,?,?)",
                     (arah, c["sstb"], c["tanggal"], c["dept_id"], c["pengrajin_id"])).lastrowid
    db.executemany("INSERT INTO dokumen_item(dokumen_id,motif_id,ket_id,jumlah,link_produksi_id) VALUES(?,?,?,?,?)",
                   [(did, x["motif_id"], x["ket_id"], x["jumlah"], x.get("link_produksi_id")) for x in c["items"]])
    return did

def save_doc(name, doc_id=None):
    arah = arah_or_404(name)
    db = get_db()
    if doc_id and not db.execute("SELECT 1 FROM dokumen WHERE id=? AND arah=?", (doc_id, arah)).fetchone():
        abort(404)
    c, errs, warns = validate_doc(db, arah, request.get_json(force=True) or {}, doc_id)
    if errs:
        return jsonify(ok=False, errors=[dict(i=i, msg=m) for i, m in errs]), 400
    if doc_id is None:
        insert_doc(db, arah, c)
    else:
        db.execute("UPDATE dokumen SET sstb=?,tanggal=?,dept_id=?,pengrajin_id=? WHERE id=?",
                   (c["sstb"], c["tanggal"], c["dept_id"], c["pengrajin_id"], doc_id))
        db.execute("DELETE FROM dokumen_item WHERE dokumen_id=?", (doc_id,))
        db.executemany("INSERT INTO dokumen_item(dokumen_id,motif_id,ket_id,jumlah,link_produksi_id) VALUES(?,?,?,?,?)",
                       [(doc_id, x["motif_id"], x["ket_id"], x["jumlah"], x.get("link_produksi_id")) for x in c["items"]])
    db.commit()
    return jsonify(ok=True, warnings=warns)

@app.post("/api/doc/<name>")
def api_doc_create(name):
    return save_doc(name)

@app.put("/api/doc/<name>/<int:doc_id>")
def api_doc_update(name, doc_id):
    return save_doc(name, doc_id)

@app.delete("/api/doc/<name>/<int:doc_id>")
def api_doc_delete(name, doc_id):
    arah = arah_or_404(name)
    db = get_db()
    db.execute("DELETE FROM dokumen WHERE id=? AND arah=?", (doc_id, arah))
    db.commit()
    return jsonify(ok=True)

@app.delete("/api/item/<name>/<int:item_id>")
def api_item_delete(name, item_id):
    arah = arah_or_404(name)
    db = get_db()
    r = db.execute("""SELECT i.dokumen_id FROM dokumen_item i JOIN dokumen d ON d.id=i.dokumen_id
                      WHERE i.id=? AND d.arah=?""", (item_id, arah)).fetchone()
    if r:
        db.execute("DELETE FROM dokumen_item WHERE id=?", (item_id,))
        if not db.execute("SELECT 1 FROM dokumen_item WHERE dokumen_id=?", (r["dokumen_id"],)).fetchone():
            db.execute("DELETE FROM dokumen WHERE id=?", (r["dokumen_id"],))
        db.commit()
    return jsonify(ok=True)

# --------------------------------------------------------------------------
# Logika stok: saldo akhir hari N = saldo awal hari N+1
# --------------------------------------------------------------------------
def stock_range(db, f, t):
    """Laporan per Jenis untuk rentang [f, t]. Saldo awal = saldo akhir hari sebelum f."""
    op = {}  # jenis -> saldo pembuka (jumlah semua motif yang tanggal mulainya sudah tercapai)
    for r in db.execute("""SELECT j.nama jenis, o.jumlah, o.tanggal_mulai FROM opening o
                           JOIN ref_motif m ON m.id=o.kode_motif_id JOIN ref_jenis j ON j.id=m.jenis_id"""):
        op.setdefault(r["jenis"], 0)
        if not r["tanggal_mulai"] or t >= r["tanggal_mulai"]:
            op[r["jenis"]] += r["jumlah"]

    def agg(arah):
        rows = db.execute("""
            SELECT j.nama jn,
                   SUM(CASE WHEN d.tanggal<? THEN i.jumlah ELSE 0 END) pre,
                   SUM(CASE WHEN d.tanggal>=? AND d.tanggal<=? THEN i.jumlah ELSE 0 END) cur
            FROM dokumen_item i JOIN dokumen d ON d.id=i.dokumen_id
            JOIN ref_motif m ON m.id=i.motif_id JOIN ref_jenis j ON j.id=m.jenis_id
            LEFT JOIN opening o ON o.kode_motif_id=m.id
            WHERE d.arah=? AND d.tanggal>=COALESCE(o.tanggal_mulai,'')
            GROUP BY j.id""", (f, f, t, arah)).fetchall()
        return {r["jn"]: (r["pre"] or 0, r["cur"] or 0) for r in rows}

    def agg_dept(arah):
        """Rincian per Dept dalam rentang [f,t] saja (bukan sepanjang histori): dari/ke dept mana."""
        rows = db.execute("""SELECT j.nama jn, dp.nama dept, SUM(i.jumlah) jml
            FROM dokumen_item i JOIN dokumen d ON d.id=i.dokumen_id
            JOIN ref_motif m ON m.id=i.motif_id JOIN ref_jenis j ON j.id=m.jenis_id
            JOIN ref_dept dp ON dp.id=d.dept_id
            WHERE d.arah=? AND d.tanggal>=? AND d.tanggal<=?
            GROUP BY j.id, dp.id ORDER BY jml DESC""", (arah, f, t)).fetchall()
        out = {}
        for r in rows:
            out.setdefault(r["jn"], []).append((r["dept"], r["jml"]))
        return out

    def dept_str(pairs):
        return ", ".join(f"{d} {fmt_id(j)}" for d, j in pairs)

    m, k = agg("M"), agg("K")
    md, kd = agg_dept("M"), agg_dept("K")
    out = []
    for j in sorted(set(op) | set(m) | set(k)):
        base = op.get(j, 0)
        mp, mc = m.get(j, (0, 0))
        kp, kc = k.get(j, (0, 0))
        awal = base + mp - kp
        out.append(dict(jenis=j, awal=awal, masuk=mc, keluar=kc, akhir=awal + mc - kc,
                        masuk_dept=dept_str(md.get(j, [])), keluar_dept=dept_str(kd.get(j, []))))
    return out

def add_total(rows, **extra):
    tot = dict(jenis="TOTAL", total=True, masuk_dept="", keluar_dept="", **extra)
    for k in ("awal", "masuk", "keluar", "akhir"):
        tot[k] = sum(r[k] for r in rows)
    return tot

def stock_data(db, f, t, mode):
    if mode == "harian":
        d0, d1 = dt.date.fromisoformat(f), dt.date.fromisoformat(t)
        if (d1 - d0).days > 92:
            d1 = d0 + dt.timedelta(days=92)
        out, d = [], d0
        while d <= d1:
            day = d.isoformat()
            rows = stock_range(db, day, day)
            rows = [r for r in rows if r["masuk"] or r["keluar"]]  # lewati jenis tanpa aktivitas pada hari itu
            for r in rows:
                r["tanggal"] = day
            out += rows
            if rows:
                out.append(add_total(rows, tanggal=day))
            d += dt.timedelta(days=1)
        return out
    rows = stock_range(db, f, t)
    return rows + ([add_total(rows)] if rows else [])

def date_args():
    today = dt.date.today().isoformat()
    f = parse_date(request.args.get("from")) or today
    t = parse_date(request.args.get("to")) or f
    if t < f:
        f, t = t, f
    return f, t

@app.get("/api/stok")
def api_stok():
    f, t = date_args()
    mode = request.args.get("mode", "ringkas")
    data = stock_data(get_db(), f, t, mode)
    return jsonify(rows=data, **{"from": f, "to": t}, mode=mode)

BY = {"tanggal": ("t.tgl_produksi", "Tanggal"), "nama": ("r_nama.nama", "Nama"), "jenis": ("rj.nama", "Jenis"),
      "motif": ("r_kode_motif.kode_motif || ' - ' || r_kode_motif.nama_motif", "Motif"),
      "ket": ("r_ket.nama", "Ket")}

def rekap_data(db, f, t, by):
    expr, label = BY.get(by, BY["tanggal"])
    _, joins = select_parts("produksi")
    sums = ",".join(f"SUM(t.{c}) {c}" for c in CATS)
    rows = db.execute(f"""SELECT COALESCE({expr},'-') grp, COUNT(*) baris, {sums}, SUM({JUMLAH_EXPR}) jumlah
                          FROM produksi t {joins} WHERE t.tgl_produksi>=? AND t.tgl_produksi<=?
                          GROUP BY grp ORDER BY grp""", (f, t)).fetchall()
    rows = [dict(r) for r in rows]
    if rows:
        tot = {"grp": "TOTAL", "total": True, "baris": sum(r["baris"] for r in rows)}
        for c in CATS + ["jumlah"]:
            tot[c] = sum(r[c] or 0 for r in rows)
        rows.append(tot)
    return label, rows

@app.get("/api/rekap")
def api_rekap():
    f, t = date_args()
    label, rows = rekap_data(get_db(), f, t, request.args.get("by", "tanggal"))
    return jsonify(rows=rows, label=label, cats=CATS)

LINE_BY = {"tanggal": ("d.tanggal", "Tanggal"), "dept": ("dp.nama", "Dept"), "jenis": ("j.nama", "Jenis"),
           "motif": ("m.kode_motif || ' - ' || m.nama_motif", "Motif"), "ket": ("k.nama", "Ket"),
           "pengrajin": ("COALESCE(pg.nama,'-')", "Nama Pengrajin")}

def rekap_lines_data(db, arah, f, t, by):
    expr, label = LINE_BY.get(by, LINE_BY["tanggal"])
    rows = db.execute(f"""SELECT COALESCE({expr},'-') grp, COUNT(*) baris, COUNT(DISTINCT d.id) sstb, SUM(i.jumlah) jumlah
        {LINE_FROM} WHERE d.arah=? AND d.tanggal>=? AND d.tanggal<=?
        GROUP BY grp ORDER BY grp""", (arah, f, t)).fetchall()
    rows = [dict(r) for r in rows]
    if rows:
        rows.append({"grp": "TOTAL", "total": True, "baris": sum(r["baris"] for r in rows),
                     "sstb": sum(r["sstb"] for r in rows), "jumlah": sum(r["jumlah"] or 0 for r in rows)})
    return label, rows

@app.get("/api/rekap/<name>")
def api_rekap_lines(name):
    arah = arah_or_404(name)
    f, t = date_args()
    label, rows = rekap_lines_data(get_db(), arah, f, t, request.args.get("by", "tanggal"))
    return jsonify(rows=rows, label=label)

# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM_FMT = '#,##0;-#,##0;"-"'

def hdr(ws, row, labels, colors):
    for i, (lab, col) in enumerate(zip(labels, colors), 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.fill = PatternFill("solid", fgColor=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 30

def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x

def send_wb(wb, name):
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

def put_cell(ws, row, col, value, kind):
    c = ws.cell(row=row, column=col, value=value)
    c.border = BORDER
    if kind == "date":
        c.number_format = "dd-mmm-yy"
    elif kind == "int":
        c.number_format = NUM_FMT
    return c

@app.get("/export/<key>")
def export_key(key):
    if key == "stok":
        return export_stok()
    if key == "rekap":
        return export_rekap(request.args.get("arah"))
    if key in ARAH:
        return export_lines(key)
    s = spec_or_404(key)
    db = get_db()
    sel, joins, w, params, order = list_query(key, request.args)
    rows = db.execute(f"SELECT {sel} FROM {s['table']} t {joins} {w} ORDER BY {order}", params).fetchall()
    wb = Workbook(); ws = wb.active; ws.title = s["title"][:31]
    hdr(ws, 1, [c["l"] for c in s["cols"]], [c.get("hc", s["color"]) for c in s["cols"]])
    for r in rows:
        ri = ws.max_row + 1
        for i, c in enumerate(s["cols"], 1):
            v = r[c["n"]]
            kind = "date" if c["t"] == "date" else "int" if (c["t"] == "int" or c.get("num")) else ""
            put_cell(ws, ri, i, dt.date.fromisoformat(v) if (kind == "date" and v) else v, kind)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths(ws, [max(12, min(34, len(c["l"]) + 6)) for c in s["cols"]])
    return send_wb(wb, f"{key}_{dt.date.today():%Y%m%d}.xlsx")

def export_lines(name):
    arah = ARAH[name]
    w, params, order = lines_query(arah, request.args)
    rows = get_db().execute(f"SELECT {LINE_SEL} {LINE_FROM} {w} ORDER BY {order}", params).fetchall()
    wb = Workbook(); ws = wb.active; ws.title = DOC_TITLE[name]
    red = {"motif", "jenis", "rumus"}
    hdr(ws, 1, [c[1] for c in LINE_COLS], ["FF0000" if c[0] in red else DOC_COLOR[name] for c in LINE_COLS])
    for r in rows:
        ri = ws.max_row + 1
        for i, (n, _, kind) in enumerate(LINE_COLS, 1):
            v = r[n]
            put_cell(ws, ri, i, dt.date.fromisoformat(v) if (kind == "date" and v) else v, kind)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths(ws, [22, 12, 16, 12, 30, 12, 12, 26, 20, 12, 30])
    return send_wb(wb, f"{name}_{dt.date.today():%Y%m%d}.xlsx")

def export_stok():
    f, t = date_args()
    mode = request.args.get("mode", "ringkas")
    rows = stock_data(get_db(), f, t, mode)
    wb = Workbook(); ws = wb.active; ws.title = "Posisi Stok"
    harian = mode == "harian"
    ncol = 9 if harian else 8
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws["A1"] = "LAPORAN POSISI STOCK GRADE"
    ws["A1"].font = Font(bold=True, size=14); ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    ws["A2"] = dt.date.fromisoformat(f) if f == t else f"{dt.date.fromisoformat(f):%d-%b-%Y} s/d {dt.date.fromisoformat(t):%d-%b-%Y}"
    if f == t:
        ws["A2"].number_format = "d-mmm-yy"
    ws["A2"].alignment = Alignment(horizontal="center")
    labels = (["Tanggal"] if harian else []) + \
             ["No", "Jenis", "Saldo Awal", "Masuk", "Dari Dept", "Keluar", "Ke Dept", "Saldo Akhir"]
    hdr(ws, 4, labels, ["4472C4"] * len(labels))
    n = 0
    for r in rows:
        if not r.get("total"):
            n += 1
        line = ([dt.date.fromisoformat(r["tanggal"])] if harian else []) + \
               ["" if r.get("total") else n, r["jenis"], r["awal"], r["masuk"], r.get("masuk_dept", ""),
                r["keluar"], r.get("keluar_dept", ""), r["akhir"]]
        ws.append(line)
        rr = ws.max_row
        numcols = {len(line) - 5, len(line) - 4, len(line) - 2, len(line)}  # Saldo Awal, Masuk, Keluar, Saldo Akhir
        for i in range(1, len(line) + 1):
            c = ws.cell(row=rr, column=i); c.border = BORDER
            if harian and i == 1:
                c.number_format = "dd-mmm-yy"
            if i in numcols:
                c.number_format = NUM_FMT
            if r.get("total"):
                c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="D9E1F2")
    widths(ws, ([12] if harian else []) + [6, 22, 12, 11, 28, 11, 28, 12])
    return send_wb(wb, f"posisi_stok_{f}_{t}.xlsx")

def export_rekap(name=None):
    f, t = date_args()
    if name in ARAH:
        arah = ARAH[name]
        label, rows = rekap_lines_data(get_db(), arah, f, t, request.args.get("by", "tanggal"))
        wb = Workbook(); ws = wb.active; ws.title = f"Rekap {DOC_TITLE[name]}"
        labels = [label, "Baris", "SSTB", "Jumlah"]
        hdr(ws, 1, labels, [DOC_COLOR[name]] * len(labels))
        for r in rows:
            ws.append([r["grp"], r["baris"], r["sstb"], r["jumlah"] or 0])
            for i in range(1, len(labels) + 1):
                c = ws.cell(row=ws.max_row, column=i); c.border = BORDER
                if i > 1:
                    c.number_format = NUM_FMT
                if r.get("total"):
                    c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="D9E1F2")
        widths(ws, [30, 8, 8, 12])
        ws.freeze_panes = "B2"
        return send_wb(wb, f"rekap_{name}_{f}_{t}.xlsx")
    label, rows = rekap_data(get_db(), f, t, request.args.get("by", "tanggal"))
    wb = Workbook(); ws = wb.active; ws.title = "Rekap Produksi"
    labels = [label, "Baris"] + [c.capitalize() for c in CATS] + ["Jumlah"]
    hdr(ws, 1, labels, ["4472C4"] * len(labels))
    for r in rows:
        ws.append([r["grp"], r["baris"]] + [r[c] or 0 for c in CATS] + [r["jumlah"] or 0])
        for i in range(1, len(labels) + 1):
            c = ws.cell(row=ws.max_row, column=i); c.border = BORDER
            if i > 1:
                c.number_format = NUM_FMT
            if r.get("total"):
                c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="D9E1F2")
    widths(ws, [30, 8] + [11] * len(CATS) + [12])
    ws.freeze_panes = "B2"
    return send_wb(wb, f"rekap_produksi_{f}_{t}.xlsx")

DOC_IN = [("sstb", "SSTB", []), ("tanggal", "Tanggal", []), ("dept", "Dept", []),
          ("kode_motif", "Kode Motif", ["kode"]), ("jumlah", "Jumlah", []), ("ket", "Ket", []),
          ("pengrajin", "Nama Pengrajin", ["pengrajin"])]

@app.get("/template/<key>")
def template_xlsx(key):
    wb = Workbook(); ws = wb.active; ws.title = "Template"
    if key in ARAH:
        hdr(ws, 1, [c[1] for c in DOC_IN], [DOC_COLOR[key]] * len(DOC_IN))
        widths(ws, [22, 12, 16, 14, 12, 26, 20])
    else:
        s = spec_or_404(key)
        cols = [c for c in s["cols"] if c["t"] != "auto"]
        hdr(ws, 1, [c["l"] for c in cols], [s["color"]] * len(cols))
        widths(ws, [max(14, len(c["l"]) + 6) for c in cols])
    return send_wb(wb, f"template_{key}.xlsx")

def norm(s):
    return re.sub(r"\s+", " ", str(s or "").strip().lower())

def read_sheet(fileobj, fields):
    """fields: [(nama, label, alias[])]. Return (rows[(no, dict)], error_json|None)."""
    if not fileobj:
        return None, (jsonify(ok=False, errors=[{"row": 0, "msg": "File belum dipilih"}]), 400)
    try:
        ws = load_workbook(fileobj, data_only=True).active
    except Exception:
        return None, (jsonify(ok=False, errors=[{"row": 0, "msg": "File bukan .xlsx yang valid"}]), 400)
    names = {}
    for n, l, al in fields:
        for a in [l, n] + al:
            names[norm(a)] = n
    hrow, colmap = None, {}
    for r in range(1, min(ws.max_row, 15) + 1):
        m = {}
        for ci in range(1, ws.max_column + 1):
            k = names.get(norm(ws.cell(row=r, column=ci).value))
            if k and k not in m.values():
                m[ci] = k
        if len(m) >= max(1, min(len(fields) // 2, 4)):
            hrow, colmap = r, m
            break
    if hrow is None:
        return None, (jsonify(ok=False, errors=[{"row": 0, "msg": "Baris header tidak dikenali. Unduh Template untuk format yang benar."}]), 400)
    rows = []
    for r in range(hrow + 1, ws.max_row + 1):
        row = {k: ws.cell(row=r, column=ci).value for ci, k in colmap.items()}
        if any(v not in (None, "") for v in row.values()):
            rows.append((r, row))
    return rows, None

@app.post("/import/<key>")
def import_xlsx(key):
    db = get_db()
    if key in ARAH:
        return import_docs(db, key)
    s = spec_or_404(key)
    cols = [c for c in s["cols"] if c["t"] != "auto"]
    rows, err = read_sheet(request.files.get("file"), [(c["n"], c["l"], c.get("alias", [])) for c in cols])
    if err:
        return err
    ok, errors = 0, []
    for rn, row in rows:
        out, errs = clean_row(db, key, row)
        if errs:
            errors.append({"row": rn, "msg": "; ".join(errs)}); continue
        cn = list(out.keys())
        db.execute(f"INSERT INTO {s['table']}({','.join(cn)}) VALUES ({','.join('?'*len(cn))})", [out[c] for c in cn])
        ok += 1
    db.commit()
    return jsonify(ok=True, inserted=ok, unit="baris", errors=errors[:200], error_count=len(errors))

def import_docs(db, name):
    arah = ARAH[name]
    rows, err = read_sheet(request.files.get("file"), DOC_IN)
    if err:
        return err
    groups, errors = {}, []
    for rn, row in rows:
        sstb = text_of(row.get("sstb"))
        if not sstb:
            errors.append({"row": rn, "msg": "SSTB kosong"}); continue
        g = groups.get(sstb)
        if g is None:
            g = groups[sstb] = dict(hdr=dict(sstb=sstb, tanggal=row.get("tanggal"), dept=row.get("dept"),
                                            pengrajin=row.get("pengrajin")), items=[], rows=[])
        elif parse_date(row.get("tanggal")) != parse_date(g["hdr"]["tanggal"]) or \
                text_of(row.get("dept")).lower() != text_of(g["hdr"]["dept"]).lower():
            errors.append({"row": rn, "msg": f"SSTB {sstb}: Tanggal/Dept berbeda dari baris pertama SSTB ini"}); continue
        g["items"].append(dict(kode_motif=row.get("kode_motif"), ket=row.get("ket"), jumlah=row.get("jumlah")))
        g["rows"].append(rn)
    n_doc = 0
    for sstb, g in groups.items():
        c, errs, _ = validate_doc(db, arah, dict(g["hdr"], items=g["items"]))
        if errs:
            for i, msg in errs:
                errors.append({"row": g["rows"][i] if i is not None else g["rows"][0], "msg": f"SSTB {sstb}: {msg}"})
            continue
        insert_doc(db, arah, c)
        n_doc += 1
    db.commit()
    return jsonify(ok=True, inserted=n_doc, unit="SSTB", errors=errors[:200], error_count=len(errors))

# --------------------------------------------------------------------------
# Backup
# --------------------------------------------------------------------------
@app.get("/backup")
def backup_download():
    os.makedirs(BACKUP_DIR, exist_ok=True)
    p = os.path.join(BACKUP_DIR, f"manual-{dt.datetime.now():%Y%m%d-%H%M%S}.db")
    src = sqlite3.connect(DB_PATH); dst = sqlite3.connect(p)
    src.backup(dst); dst.close(); src.close()
    return send_file(p, as_attachment=True, download_name=os.path.basename(p))

def daily_backup(keep=30):
    os.makedirs(BACKUP_DIR, exist_ok=True)
    p = os.path.join(BACKUP_DIR, f"inventory-{dt.date.today():%Y%m%d}.db")
    if not os.path.exists(p) and os.path.exists(DB_PATH):
        src = sqlite3.connect(DB_PATH); dst = sqlite3.connect(p)
        src.backup(dst); dst.close(); src.close()
    files = sorted(f for f in os.listdir(BACKUP_DIR) if f.startswith("inventory-"))
    for old in files[:-keep]:
        os.remove(os.path.join(BACKUP_DIR, old))

# --------------------------------------------------------------------------
# Halaman
# --------------------------------------------------------------------------
PAGES = [("db", "1. DB"), ("masuk", "2. Barang Masuk"), ("produksi", "3. Produksi"),
         ("keluar", "4. Barang Keluar"), ("stok", "5. Laporan Stok")]

@app.context_processor
def inject():
    return dict(pages=PAGES, current=request.path.strip("/").split("/")[0])

@app.get("/")
def home():
    return redirect("/stok")

@app.get("/db")
def page_db():
    tabs = [(k, SPECS[k]["title"]) for k in ("motif", "jenis", "opening", "ket", "dept", "ketprod", "pengrajin", "karyawan")]
    return render_template("db.html", tabs=tabs, title="DB (Master Data)")

@app.get("/<key>")
def page_main(key):
    if key == "produksi":
        return render_template("produksi.html", title="Produksi")
    if key in ARAH:
        return render_template("doc.html", arah=key, title=DOC_TITLE[key])
    if key == "stok":
        return render_template("stok.html", title="Laporan Posisi Stok")
    if key == "rekap":
        return redirect("/produksi")
    abort(404)

def lan_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM); s.connect(("10.255.255.255", 1))
        ip = s.getsockname()[0]; s.close(); return ip
    except Exception:
        return "127.0.0.1"

init_db()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    daily_backup()
    print(f"\n  Stok Grade berjalan.\n  Buka di komputer ini : http://localhost:{port}"
          f"\n  Dari komputer lain   : http://{lan_ip()}:{port}\n  (Tutup jendela ini untuk menghentikan)\n")
    threading.Timer(1.2, lambda: webbrowser.open(f"http://localhost:{port}")).start()
    try:
        from waitress import serve
        serve(app, host="0.0.0.0", port=port, threads=4)
    except ImportError:
        app.run(host="0.0.0.0", port=port)
